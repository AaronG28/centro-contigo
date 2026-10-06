"""Exportaciones a Excel y PDF."""

from datetime import datetime

from PySide6.QtCore import QMarginsF
from PySide6.QtGui import (QFont, QPageLayout, QPageSize, QPdfWriter,
                           QTextDocument)

from ..config import CAMPOS_REAL, EXPORTABLES
from ..db.connection import db
from ..sesion import USUARIO_ACTUAL
from .resumen import calcular_resumen


def exportar_excel(path, mes):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="0071E3")
    money = '"S/" #,##0.00'

    def estilo_cabecera(ws, ncols):
        for j in range(1, ncols + 1):
            c = ws.cell(row=1, column=j)
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(horizontal="center")

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    d = calcular_resumen(mes)
    ws.append([f"Resumen — {'todos los periodos' if mes == 'Todos' else mes}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Ingresos", d["ti"]]); ws.append(["Egresos", d["te"]])
    ws.append(["Balance", d["ti"] - d["te"]]); ws.append(["Pendiente de cobro", d["pend"]])
    for r in range(3, 7):
        ws.cell(row=r, column=1).font = Font(bold=True)
        ws.cell(row=r, column=2).number_format = money
    ws.append([])
    ws.append(["Cuenta", "Ingresos", "Egresos", "Saldo"])
    estilo = ws.max_row
    for j in range(1, 5):
        ws.cell(row=estilo, column=j).font = head_font
        ws.cell(row=estilo, column=j).fill = head_fill
    for f in d["filas"]:
        ws.append(list(f))
        for j in (2, 3, 4):
            ws.cell(row=ws.max_row, column=j).number_format = money
    for col, w in zip("ABCD", (28, 16, 16, 16)):
        ws.column_dimensions[col].width = w

    with db() as c:
        for nombre, tabla in EXPORTABLES.items():
            hoja = wb.create_sheet(nombre)
            cur = c.execute(f"SELECT * FROM {tabla} ORDER BY id")
            cols = [x[0] for x in cur.description]
            hoja.append(cols)
            for row in cur:
                hoja.append(list(row))
            estilo_cabecera(hoja, len(cols))
            hoja.freeze_panes = "A2"
            hoja.auto_filter.ref = hoja.dimensions
            for j, col in enumerate(cols, 1):
                letra = get_column_letter(j)
                ancho = max([len(str(col))] + [
                    len(str(hoja.cell(row=r, column=j).value or ""))
                    for r in range(2, min(hoja.max_row, 200) + 1)])
                hoja.column_dimensions[letra].width = min(max(ancho + 2, 10), 45)
                if col in CAMPOS_REAL:
                    for r in range(2, hoja.max_row + 1):
                        hoja.cell(row=r, column=j).number_format = money
    wb.save(path)


def _html_tabla(titulo, heads, rows):
    h = [f"<h3>{titulo}</h3><table width='100%' cellspacing='0' cellpadding='4' "
         "border='1' style='border-color:#d2d2d7'><tr style='background:#f0f4fa'>"]
    h += [f"<th align='left'>{x}</th>" for x in heads]
    h.append("</tr>")
    for r in rows:
        h.append("<tr>")
        for j, v in enumerate(r):
            txt = f"S/ {v:,.2f}" if isinstance(v, float) else str(v)
            h.append(f"<td align='{'right' if j else 'left'}'>{txt}</td>")
        h.append("</tr>")
    if not rows:
        h.append(f"<tr><td colspan='{len(heads)}'><i>Sin datos</i></td></tr>")
    h.append("</table>")
    return "".join(h)


def exportar_pdf(path, mes):
    d = calcular_resumen(mes)
    per = "Todos los periodos" if mes == "Todos" else mes
    usuario = USUARIO_ACTUAL.get("nombre", "—")
    html = f"""
    <h1 style='color:#0071e3'>Centro Contigo — Informe de resumen</h1>
    <p>Periodo: <b>{per}</b> · Generado por <b>{usuario}</b>
       el {datetime.now():%d/%m/%Y %H:%M}</p>
    <table width='100%' cellpadding='6'><tr>
      <td><b>Ingresos</b><br><span style='color:#34c759;font-size:16pt'>
      S/ {d['ti']:,.2f}</span></td>
      <td><b>Egresos</b><br><span style='color:#ff3b30;font-size:16pt'>
      S/ {d['te']:,.2f}</span></td>
      <td><b>Balance</b><br><span style='color:#0071e3;font-size:16pt'>
      S/ {d['ti'] - d['te']:,.2f}</span></td>
      <td><b>Por cobrar</b><br><span style='font-size:16pt'>
      S/ {d['pend']:,.2f}</span></td>
    </tr></table>
    {_html_tabla('Movimiento por cuenta', ['Cuenta', 'Ingresos', 'Egresos', 'Saldo'],
                 d['filas'])}
    {_html_tabla('Ingresos por tipo', ['Tipo', 'N°', 'Monto'], d['tip'])}
    {_html_tabla('Ingresos por terapeuta', ['Terapeuta', 'N°', 'Monto'], d['ter'])}
    {_html_tabla('Egresos por tipo', ['Tipo', 'N°', 'Monto'], d['tegr'])}
    {_html_tabla('Flujo de caja — últimos 12 meses', ['Mes', 'Ingresos', 'Egresos'],
                 list(zip(d['meses'], d['s_ing'], d['s_egr'])))}
    """
    doc = QTextDocument()
    doc.setDefaultFont(QFont("Helvetica Neue", 9))
    doc.setHtml(html)
    w = QPdfWriter(path)
    w.setPageSize(QPageSize(QPageSize.A4))
    w.setPageMargins(QMarginsF(15, 15, 15, 15), QPageLayout.Millimeter)
    doc.print_(w)
